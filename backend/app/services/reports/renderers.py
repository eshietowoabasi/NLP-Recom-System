"""PDF (reportlab) and DOCX (python-docx) renderers for ReportContent."""
import io
from xml.sax.saxutils import escape

from ..preprocessing.normalize import normalize_unicode


# --- DOCX -------------------------------------------------------------------

def render_docx(content) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    document = docx.Document()
    document.styles["Normal"].font.size = Pt(10)
    heading = document.add_heading(content.title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(content.subtitle)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for label, value in content.meta:
        p = document.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)

    for block in content.blocks:
        if block.kind == "heading":
            document.add_heading(block.text, level=1)
        elif block.kind == "subheading":
            document.add_heading(block.text, level=2)
        elif block.kind == "paragraph":
            document.add_paragraph(block.text)
        elif block.kind == "bullets":
            for item in block.items:
                document.add_paragraph(item, style="List Bullet")
        elif block.kind == "table":
            table = document.add_table(rows=1, cols=len(block.headers))
            table.style = "Light Grid Accent 1"
            for cell, header in zip(table.rows[0].cells, block.headers):
                cell.text = header
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in block.rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = str(value)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# --- PDF --------------------------------------------------------------------

def _pdf_text(text) -> str:
    """Escape markup and keep to characters the built-in PDF fonts can draw."""
    text = normalize_unicode(str(text))
    text = text.encode("latin-1", "replace").decode("latin-1")
    return escape(text)


def render_pdf(content) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    for name in ("Heading1", "Heading2", "Heading3"):
        styles[name].keepWithNext = 1  # never strand a heading at the bottom of a page
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=9.5, leading=12.5)
    cell = ParagraphStyle("cell", parent=body, fontSize=8, leading=10)
    cell_bold = ParagraphStyle("cellb", parent=cell, fontName="Helvetica-Bold")

    story = [
        Paragraph(_pdf_text(content.title), styles["Title"]),
        Paragraph(_pdf_text(content.subtitle), ParagraphStyle("sub", parent=styles["Heading2"], alignment=1)),
        Spacer(1, 0.3 * cm),
    ]
    for label, value in content.meta:
        story.append(Paragraph(f"<b>{_pdf_text(label)}:</b> {_pdf_text(value)}", body))
    story.append(Spacer(1, 0.4 * cm))

    width = A4[0] - 4 * cm
    for block in content.blocks:
        if block.kind == "heading":
            story.append(Paragraph(_pdf_text(block.text), styles["Heading1"]))
        elif block.kind == "subheading":
            story.append(Paragraph(_pdf_text(block.text), styles["Heading3"]))
        elif block.kind == "paragraph":
            story.append(Paragraph(_pdf_text(block.text), body))
            story.append(Spacer(1, 0.15 * cm))
        elif block.kind == "bullets" and block.items:
            story.append(ListFlowable(
                [ListItem(Paragraph(_pdf_text(i), body), leftIndent=12) for i in block.items],
                bulletType="bullet", start="\u2022", leftIndent=12,
            ))
        elif block.kind == "table":
            data = [[Paragraph(_pdf_text(h), cell_bold) for h in block.headers]]
            data += [[Paragraph(_pdf_text(v), cell) for v in row] for row in block.rows]
            weights = block.widths or [1] * len(block.headers)
            table = Table(data, colWidths=[width * w / sum(weights) for w in weights], repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dbe5f1")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa5b1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(table)
            story.append(Spacer(1, 0.3 * cm))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(2 * cm, 1.2 * cm, _pdf_text(f"{content.title} - {content.subtitle}")[:110])
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                      bottomMargin=2 * cm, title=content.title, author="NLP-RS").build(
        story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


RENDERERS = {"pdf": render_pdf, "docx": render_docx}
