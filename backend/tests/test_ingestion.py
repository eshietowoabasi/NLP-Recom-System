import pytest

from app.models import FileType
from app.services.ingestion import (
    FileTooLargeError,
    ParseError,
    UnsupportedFileError,
    parse_document,
    validate_upload,
)
from app.utils.errors import ValidationError

from .fixtures import JOB_AD, make_docx, make_pdf, make_scanned_pdf

MB = 1024 * 1024


class TestValidation:
    def test_accepts_supported_formats(self):
        assert validate_upload("a.pdf", make_pdf([JOB_AD]), 25 * MB) is FileType.PDF
        assert validate_upload("a.DOCX", make_docx([JOB_AD]), 25 * MB) is FileType.DOCX
        assert validate_upload("a.txt", JOB_AD.encode(), 25 * MB) is FileType.TXT

    @pytest.mark.parametrize("name", ["a.doc", "a.exe", "a.html", "noextension"])
    def test_rejects_unsupported_extensions(self, name):
        with pytest.raises(UnsupportedFileError):
            validate_upload(name, b"hello", 25 * MB)

    @pytest.mark.parametrize(
        "name, data",
        [
            ("fake.pdf", b"MZ\x90\x00 this is an executable"),
            ("fake.docx", b"PK\x03\x04 not really a zip"),
            ("fake.txt", b"\x89PNG\r\n\x1a\n\x00\x00\x00binary"),
        ],
        ids=["exe-as-pdf", "bad-zip-as-docx", "png-as-txt"],
    )
    def test_rejects_content_that_does_not_match_extension(self, name, data):
        with pytest.raises(UnsupportedFileError):
            validate_upload(name, data, 25 * MB)

    def test_rejects_plain_zip_named_docx(self):
        import io
        import zipfile

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("hello.txt", "hi")
        with pytest.raises(UnsupportedFileError):
            validate_upload("x.docx", buf.getvalue(), 25 * MB)

    def test_size_limit(self):
        with pytest.raises(FileTooLargeError):
            validate_upload("a.txt", b"a" * (MB + 1), MB)

    def test_empty_file(self):
        with pytest.raises(ValidationError):
            validate_upload("a.txt", b"", MB)

    def test_cp1252_text_accepted(self):
        assert validate_upload("a.txt", "Café – résumé".encode("cp1252"), MB) is FileType.TXT


class TestParsers:
    def test_txt_blocks_and_encoding(self):
        data = "First paragraph\nstill first.\n\nSecond – café.".encode("cp1252")
        parsed = parse_document(data, FileType.TXT)
        assert parsed.blocks == ["First paragraph\nstill first.", "Second – café."]
        assert parsed.word_count == 6

    def test_txt_utf8_bom(self):
        parsed = parse_document("﻿Hello world".encode("utf-8"), FileType.TXT)
        assert parsed.blocks == ["Hello world"]

    def test_docx_body_and_tables_in_order_without_header_footer(self):
        data = make_docx(
            ["Intro paragraph.", "Skills required below."],
            table=[["Skill", "Level"], ["Python", "Advanced"]],
            header="CONFIDENTIAL RUNNING HEADER",
        )
        parsed = parse_document(data, FileType.DOCX)
        assert parsed.blocks == ["Intro paragraph.", "Skills required below.", "Skill | Level", "Python | Advanced"]
        assert "CONFIDENTIAL" not in parsed.text and "Page 1" not in parsed.text

    def test_docx_merged_cells_not_duplicated(self):
        import io

        import docx

        document = docx.Document()
        table = document.add_table(rows=1, cols=3)
        merged = table.cell(0, 0).merge(table.cell(0, 1))
        merged.text = "Merged"
        table.cell(0, 2).text = "Other"
        buf = io.BytesIO()
        document.save(buf)
        assert parse_document(buf.getvalue(), FileType.DOCX).blocks == ["Merged | Other"]

    def test_docx_keeps_tab_so_toc_lines_can_be_detected(self):
        parsed = parse_document(make_docx(["Body"], toc_entries=[("PREAMBLE", 5)]), FileType.DOCX)
        assert "PREAMBLE\t5" in parsed.blocks

    def test_corrupt_docx_raises_parse_error(self):
        import io
        import zipfile

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("word/document.xml", "<not-xml")
        with pytest.raises(ParseError):
            parse_document(buf.getvalue(), FileType.DOCX)

    def test_pdf_text_pages_and_furniture_removed(self):
        pages = [JOB_AD, "Second page about data engineering.", "Third page on DevSecOps practice."]
        parsed = parse_document(make_pdf(pages, header="ACME Careers Portal"), FileType.PDF)
        assert parsed.page_count == 3
        assert "Kubernetes" in parsed.text and "DevSecOps" in parsed.text
        assert "ACME Careers Portal" not in parsed.text
        assert "Page 2 of 3" not in parsed.text

    def test_pdf_dehyphenates_line_breaks(self):
        parsed = parse_document(make_pdf(["Cyber secu-\nrity and informa-\ntion systems"]), FileType.PDF)
        assert "security" in parsed.text and "information" in parsed.text

    def test_scanned_pdf_fails_with_clear_message(self):
        with pytest.raises(ParseError, match="scanned"):
            parse_document(make_scanned_pdf(), FileType.PDF)

    def test_corrupt_pdf_raises_parse_error(self):
        with pytest.raises(ParseError):
            parse_document(b"%PDF-1.7\n garbage garbage", FileType.PDF)

    def test_whitespace_only_txt_raises(self):
        with pytest.raises(ParseError):
            parse_document(b"   \n\n  ", FileType.TXT)
