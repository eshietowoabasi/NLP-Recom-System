import io

import pytest

from app.models import AuditLog, Role
from app.schemas.mapping import normalise_course_code, validate_mapping
from app.utils.errors import ValidationError

MAPPING = {
    "course_code": "csc 413",
    "course_title": "Cloud Infrastructure and DevOps",
    "credit_units": 3,
    "prerequisites": ["CSC 201", "csc301"],
    "learning_outcomes": ["Deploy containerised services", "Automate infrastructure with code"],
}


def accepted_rec(client, session_id, overlap="No Significant Overlap"):
    rec = next(r for r in client.get(f"/api/sessions/{session_id}/recommendations").json["items"]
               if r["overlap_status"] == overlap)
    body = {"decision": "Accepted", "notes": "Justified by local demand"}
    assert client.patch(f"/api/recommendations/{rec['rec_id']}/decision", json=body).status_code == 200
    return rec["rec_id"]


# --- mapping validation -----------------------------------------------------

@pytest.mark.parametrize("raw, expected", [("csc413", "CSC 413"), (" CYB 401l ", "CYB 401L"), ("SEN312", "SEN 312")])
def test_course_code_normalised(raw, expected):
    assert normalise_course_code(raw) == expected


@pytest.mark.parametrize(
    "override, field",
    [
        ({"course_code": "Cloud101"}, "course_code"),
        ({"course_title": " "}, "course_title"),
        ({"credit_units": 4}, "credit_units"),
        ({"credit_units": True}, "credit_units"),
        ({"learning_outcomes": []}, "learning_outcomes"),
        ({"learning_outcomes": "one"}, "learning_outcomes"),
        ({"prerequisites": ["XYZ"]}, "prerequisites"),
        ({"prerequisites": ["CSC 413"]}, "prerequisites"),
    ],
)
def test_mapping_validation_errors(override, field):
    with pytest.raises(ValidationError) as exc:
        validate_mapping({**MAPPING, **override})
    assert field in exc.value.details


def test_mapping_validation_normalises():
    data = validate_mapping(MAPPING)
    assert data["course_code"] == "CSC 413" and data["prerequisites"] == ["CSC 201", "CSC 301"]


# --- mapping API --------------------------------------------------------------

def test_mapping_lifecycle(client, db, completed_session):
    items = client.get(f"/api/sessions/{completed_session}/recommendations").json["items"]
    pending = items[0]["rec_id"]
    resp = client.post(f"/api/recommendations/{pending}/mapping", json=MAPPING)
    assert resp.status_code == 409 and resp.json["error"]["code"] == "NOT_ACCEPTED"

    rec_id = accepted_rec(client, completed_session)
    resp = client.post(f"/api/recommendations/{rec_id}/mapping", json=MAPPING)
    assert resp.status_code == 201, resp.json
    mapping = resp.json["mapping"]
    assert mapping["course_code"] == "CSC 413" and mapping["rec_id"] == rec_id and mapping["topic_title"]

    # Course codes are unique within a session.
    dup_rec = accepted_rec(client, completed_session, overlap="Potential Duplicate")
    resp = client.post(f"/api/recommendations/{dup_rec}/mapping", json=MAPPING)
    assert resp.status_code == 409 and resp.json["error"]["code"] == "COURSE_CODE_IN_USE"
    assert client.post(f"/api/recommendations/{dup_rec}/mapping", json={**MAPPING, "course_code": "CYB 401"}).status_code == 201

    # A mapped recommendation cannot be un-accepted.
    resp = client.patch(f"/api/recommendations/{rec_id}/decision", json={"decision": "Rejected"})
    assert resp.status_code == 409 and resp.json["error"]["code"] == "MAPPING_EXISTS"

    updated = client.put(f"/api/mappings/{mapping['map_id']}", json={**MAPPING, "credit_units": 2, "course_title": "Cloud DevOps"})
    assert updated.status_code == 200 and updated.json["mapping"]["credit_units"] == 2
    assert client.put(f"/api/mappings/{mapping['map_id']}", json={**MAPPING, "course_code": "CYB 401"}).status_code == 409
    assert client.get(f"/api/recommendations/{rec_id}/mapping").json["items"][0]["course_title"] == "Cloud DevOps"
    assert [m["course_code"] for m in client.get(f"/api/sessions/{completed_session}/mappings").json["items"]] == ["CSC 413", "CYB 401"]

    assert client.delete(f"/api/mappings/{mapping['map_id']}").status_code == 200
    assert client.get(f"/api/mappings/{mapping['map_id']}").status_code == 404
    assert client.patch(f"/api/recommendations/{rec_id}/decision", json={"decision": "Rejected"}).status_code == 200
    actions = {a.action_type for a in db.session.query(AuditLog)}
    assert {"MAPPING_CREATE", "MAPPING_UPDATE", "MAPPING_DELETE"} <= actions


def test_mapping_permissions(client, completed_session, make_user, login):
    rec_id = accepted_rec(client, completed_session)
    map_id = client.post(f"/api/recommendations/{rec_id}/mapping", json=MAPPING).json["mapping"]["map_id"]
    client.post("/api/auth/logout")
    make_user("viewer", role=Role.VIEWER)
    login("viewer")
    assert client.get(f"/api/mappings/{map_id}").status_code == 200
    assert client.put(f"/api/mappings/{map_id}", json=MAPPING).status_code == 403
    client.post("/api/auth/logout")
    make_user("other", role=Role.PLANNER)
    login("other")
    assert client.delete(f"/api/mappings/{map_id}").status_code == 403


def test_rename_recommendation(client, completed_session):
    rec_id = client.get(f"/api/sessions/{completed_session}/recommendations").json["items"][0]["rec_id"]
    resp = client.patch(f"/api/recommendations/{rec_id}", json={"topic_title": "  Cloud Engineering Practice "})
    assert resp.status_code == 200 and resp.json["recommendation"]["topic_title"] == "Cloud Engineering Practice"
    assert client.patch(f"/api/recommendations/{rec_id}", json={"topic_title": ""}).status_code == 422
    assert client.patch(f"/api/recommendations/{rec_id}", json={}).status_code == 422


# --- reports ------------------------------------------------------------------

@pytest.mark.parametrize("fmt, magic", [("pdf", b"%PDF-"), ("docx", b"PK\x03\x04")])
def test_generate_and_download_report(client, db, completed_session, fmt, magic):
    rec_id = accepted_rec(client, completed_session)
    client.post(f"/api/recommendations/{rec_id}/mapping", json=MAPPING)

    resp = client.post(f"/api/sessions/{completed_session}/reports", json={"format": fmt})
    assert resp.status_code == 201, resp.json
    report = resp.json["report"]
    assert report["format"] == fmt and report["file_size"] > 1000 and report["session_name"] == "review"

    download = client.get(f"/api/reports/{report['report_id']}/download")
    assert download.status_code == 200 and download.data.startswith(magic)
    assert f".{fmt}" in download.headers["Content-Disposition"]

    text = _extract_text(download.data, fmt)
    for expected in ("Curriculum Recommendation Report", "Session configuration", "Ranked recommendations",
                     "Semantic overlap", "CSC 413", "Deploy containerised services", "Hashing embeddings"):
        assert expected in text, expected
    assert client.get("/api/reports").json["pagination"]["total"] == 1
    assert client.get(f"/api/reports?session_id={completed_session}").json["items"][0]["report_id"] == report["report_id"]
    assert db.session.query(AuditLog).filter_by(action_type="REPORT_GENERATE").count() == 1


def _extract_text(data, fmt):
    if fmt == "pdf":
        import pymupdf

        with pymupdf.open(stream=data, filetype="pdf") as pdf:
            return " ".join(" ".join(page.get_text().split()) for page in pdf)
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    parts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return " ".join(parts)


def test_report_errors(client, completed_session, make_user, login):
    assert client.post(f"/api/sessions/{completed_session}/reports", json={"format": "html"}).status_code == 422
    assert client.get("/api/reports/999").status_code == 404
    client.post("/api/auth/logout")
    make_user("viewer", role=Role.VIEWER)
    login("viewer")
    assert client.post(f"/api/sessions/{completed_session}/reports", json={"format": "pdf"}).status_code == 403


def test_missing_report_file_returns_404(client, db, completed_session):
    import os

    from app.models import Report

    report_id = client.post(f"/api/sessions/{completed_session}/reports", json={}).json["report"]["report_id"]
    os.remove(db.session.get(Report, report_id).file_path)
    assert client.get(f"/api/reports/{report_id}/download").status_code == 404


def test_pdf_handles_markup_and_unicode():
    from app.services.reports import Block, ReportContent, render_pdf

    content = ReportContent("T", "S <b>&", [("Who", "Ọlá")], [
        Block("paragraph", "5 < 6 & “quotes” – dash ≤"),
        Block("table", headers=["a"], rows=[["<script>"]]),
        Block("bullets", items=[]),
    ])
    assert render_pdf(content).startswith(b"%PDF-")


# --- admin ----------------------------------------------------------------------

@pytest.fixture()
def admin(make_user, login):
    user = make_user("admin", role=Role.ADMIN)
    login("admin")
    return user


def test_admin_user_management(client, admin, login):
    resp = client.post("/api/admin/users", json={"username": "ada", "email": "Ada@Uniuyo.edu.ng",
                                                  "password": "long-enough", "role": "Curriculum Planner"})
    assert resp.status_code == 201 and resp.json["user"]["email"] == "ada@uniuyo.edu.ng"
    user_id = resp.json["user"]["user_id"]
    assert client.post("/api/admin/users", json={"username": "ada", "email": "x@y.ng", "password": "long-enough"}).status_code == 409
    bad = client.post("/api/admin/users", json={"username": "a", "email": "nope", "password": "short", "role": "God"})
    assert bad.status_code == 422 and set(bad.json["error"]["details"]) == {"username", "email", "password", "role"}

    assert client.get("/api/admin/users?q=ad").json["pagination"]["total"] == 2
    assert client.patch(f"/api/admin/users/{user_id}", json={"role": "Viewer", "is_active": False}).json["user"]["role"] == "Viewer"
    assert client.patch(f"/api/admin/users/{user_id}", json={}).status_code == 422
    assert client.patch(f"/api/admin/users/{user_id}", json={"password": "new-password", "is_active": True}).status_code == 200

    # Admins cannot lock themselves out.
    assert client.patch(f"/api/admin/users/{admin.user_id}", json={"role": "Viewer"}).status_code == 422
    assert client.patch(f"/api/admin/users/{admin.user_id}", json={"is_active": False}).status_code == 422
    assert client.get("/api/auth/me").json["user"]["role"] == "Admin"

    client.post("/api/auth/logout")
    assert login("ada", "new-password").status_code == 200


def test_admin_endpoints_forbidden_for_planners(client, make_user, login):
    make_user("planner", role=Role.PLANNER)
    login("planner")
    for method, url in (("get", "/api/admin/users"), ("post", "/api/admin/users"), ("get", "/api/admin/audit")):
        assert getattr(client, method)(url).status_code == 403


def test_audit_log_filters(client, admin, make_user):
    make_user("someone")
    client.post("/api/auth/login", json={"username": "ghost", "password": "x"})  # LOGIN_FAILED while logged in as admin
    log = client.get("/api/admin/audit").json
    assert log["pagination"]["total"] >= 2 and log["items"][0]["username"] in ("admin", None)
    logins = client.get("/api/admin/audit?action_type=login").json["items"]
    assert logins and all(e["action_type"] == "LOGIN" for e in logins)
    assert client.get(f"/api/admin/audit?user_id={admin.user_id}").json["items"]
    assert client.get("/api/admin/audit?from=2000-01-01&to=2999-01-01").json["pagination"]["total"] >= 2
    assert client.get("/api/admin/audit?from=yesterday").status_code == 422


# --- dashboard ------------------------------------------------------------------

def test_dashboard_summary(client, completed_session):
    data = client.get("/api/dashboard/summary").json
    assert data["documents"]["by_category"]["Job Market Data"] == 3
    assert data["documents"]["by_category"]["NUC Core Reference"] == 1
    assert data["sessions"]["by_status"]["Completed"] == 1
    assert data["core_reference_available"] is True
    assert data["my_pending_reviews"] > 0
    assert data["recent_sessions"][0]["session_id"] == completed_session
