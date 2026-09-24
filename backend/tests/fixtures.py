"""Synthetic test documents generated at runtime (no real corpus files in the repo)."""
import io

import docx
import pymupdf as fitz


def make_docx(paragraphs, table=None, header=None, toc_entries=None):
    document = docx.Document()
    if header:
        document.sections[0].header.paragraphs[0].text = header
        document.sections[0].footer.paragraphs[0].text = "Page 1"
    if toc_entries:
        document.add_paragraph("TABLE OF CONTENTS")
        for entry, page in toc_entries:
            document.add_paragraph(f"{entry}\t{page}")
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_pdf(pages, header=None, footer_page_numbers=True):
    pdf = fitz.open()
    for number, body in enumerate(pages, start=1):
        page = pdf.new_page()
        if header:
            page.insert_text((72, 40), header)
        page.insert_textbox(fitz.Rect(72, 80, 520, 760), body, fontsize=11)
        if footer_page_numbers:
            page.insert_text((280, 810), f"Page {number} of {len(pages)}")
    data = pdf.tobytes()
    pdf.close()
    return data


def make_scanned_pdf():
    """A PDF whose only content is an image, i.e. no extractable text."""
    pdf = fitz.open()
    page = pdf.new_page()
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 10, 10), False)
    page.insert_image(fitz.Rect(0, 0, 100, 100), pixmap=pix)
    data = pdf.tobytes()
    pdf.close()
    return data


JOB_AD = (
    "We are hiring a Cloud Security Engineer with experience in AWS, Kubernetes and "
    "Python. The candidate will design secure data pipelines and automate incident "
    "response. Knowledge of machine learning is an advantage."
)
